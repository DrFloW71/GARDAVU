from collections import Counter
from datetime import date
from fractions import Fraction
from tkinter import ttk

import pytest
from openpyxl import load_workbook

from gardavu.domain import Doctor, Plan, Record, Slot, Span
from gardavu.exports import export_excel
from gardavu.solver import solve
from gardavu.storage import Project
from gardavu.targets import calculate_targets, statistics
from gardavu.ui import App, Form
from gardavu.validation import validate

pytestmark = pytest.mark.gui


def weighted_slots():
    weekdays = ['2027-01-04', '2027-01-11', '2027-01-18', '2027-01-25', '2027-02-01', '2027-02-08']
    saturdays = ['2027-02-20', '2027-02-27', '2027-03-06']
    return [Slot(day, day, 'SEMAINE', '19:00', '24:00') for day in weekdays] + [
        Slot(day, day, 'SAMEDI', '14:00', '24:00') for day in saturdays]


@pytest.mark.parametrize('weight,expected', [(2, {'A': 6, 'B': 3}), (3, {'A': 6, 'B': 2})])
def test_multiple_weight_is_used_by_solver(rules, weight, expected):
    doctors = [Doctor('Pondéré', id='A', coefficient=weight), Doctor('Normal', id='B')]
    slots = weighted_slots()[:sum(expected.values())]
    plan = solve(Plan(slots[0].day, slots[-1].day, slots), doctors, rules, [])
    assert plan.solver_status in ('OPTIMAL', 'FEASIBLE'), plan.metadata
    assert Counter(d for ids in plan.assignments.values() for d in ids) == expected
    assert validate(plan, doctors, rules, []).ok
    if weight == 2:
        rows = {row['id']: row for row in statistics(plan, doctors)}
        assert rows['A']['categories'] == {'SEMAINE': 4, 'SAMEDI': 2}
        assert rows['B']['categories'] == {'SEMAINE': 2, 'SAMEDI': 1}


def test_weight_targets_remain_proportional_by_month_category_and_holiday(rules):
    doctors = [Doctor('Pondéré', id='A', coefficient=2), Doctor('Normal', id='B')]
    slots = weighted_slots() + [Slot('noel', '2027-12-25', 'FERIE_JOUR', '08:00', '19:00', holidays=('NOEL',))]
    targets = calculate_targets(slots, doctors)
    for key in {key for ident, key in targets}:
        assert targets['A', key] == 2 * targets['B', key]
    doctors[0].suspensions = [Span('2027-01-01', '2027-01-31')]
    targets = calculate_targets(slots, doctors)
    assert targets['A', 'MOIS:2027-01'] == 0
    assert targets['B', 'MOIS:2027-01'] == 4
    assert targets['A', 'MOIS:2027-02'] == 2 * targets['B', 'MOIS:2027-02']


def test_weight_never_creates_duplicate_doctor_or_relaxes_spacing(rules):
    weighted = Doctor('Pondéré', id='A', coefficient=2)
    slot = Slot('now', '2027-02-09', 'SEMAINE', '19:00', '24:00', required=2)
    impossible = solve(Plan(slot.day, slot.day, [slot]), [weighted], rules, [])
    assert impossible.solver_status == 'INFEASIBLE'
    doctors = [weighted, Doctor('Normal', id='B')]
    filled = solve(Plan(slot.day, slot.day, [slot]), doctors, rules, [])
    assert set(filled.assignments[slot.id]) == {'A', 'B'}
    single = Slot('now', slot.day, slot.category, slot.start, slot.end)
    history = [Record(Slot('previous', '2027-02-04', 'SEMAINE', None, None), 'A')]
    result = solve(Plan(single.day, single.day, [single]), doctors, rules, history)
    assert result.assignments[single.id] == ['B']
    assert validate(result, doctors, rules, history).ok


def test_weight_button_persists_and_updates_targets_without_changing_saved_roster(tmp_path, rules):
    app = App()
    app.withdraw()
    path = tmp_path / 'Pondération.gardavu'
    try:
        project = Project(path, create=True)
        project.save_rules(rules)
        project.save_doctor(Doctor('Pondéré', id='A'))
        project.save_doctor(Doctor('Normal', id='B'))
        app.set_project(project)
        slots = weighted_slots()
        app.plan = solve(Plan(slots[0].day, slots[-1].day, slots), project.doctors(), rules, [])
        project.save_draft(app.plan)
        previous_assignments = {key: list(ids) for key, ids in app.plan.assignments.items()}
        app.refresh()
        app.doctor_tree.selection_set('A')
        buttons = [w for frame in app.tabs['Médecins et absences'].winfo_children()
                   for w in frame.winfo_children() if isinstance(w, ttk.Button)]
        next(w for w in buttons if w.cget('text') == 'Pondérer le médecin').invoke()
        dialog = next(w for w in app.winfo_children() if isinstance(w, Form))
        dialog.values['coefficient'].delete(0, 'end')
        dialog.values['coefficient'].insert(0, '2')
        dialog.accept.__wrapped__(dialog)
        assert len(project.doctors()) == 2
        assert next(d for d in project.doctors() if d.id == 'A').coefficient == 2
        assert app.doctor_tree.item('A', 'values')[1] == '×2'
        assert app.equity_tree.item('A', 'values')[1:3] == ('×2', '6.000')
        assert app.plan.assignments == previous_assignments
        assert app.plan.solver_status != 'OPTIMAL' and app.plan.metadata['weights_changed']
        assert not app.plan.metadata['stages']
        assert 'régénérer' in app.plan_status.cget('text')
        assert 'Pondération modifiée' in app.diagnostic_text()
    finally:
        app.close()
    with Project(path) as reopened:
        assert next(d for d in reopened.doctors() if d.id == 'A').coefficient == 2
        targets = calculate_targets(slots, reopened.doctors())
        assert targets['A', 'TOTAL'] == 6 and targets['B', 'TOTAL'] == 3


@pytest.mark.parametrize('value', ['0', '-2', 'nan', 'inf'])
def test_weight_dialog_rejects_nonpositive_or_nonfinite_value(tmp_path, rules, value):
    app = App()
    app.withdraw()
    try:
        project = Project(tmp_path / 'Limites.gardavu', create=True)
        project.save_rules(rules)
        project.save_doctor(Doctor('Normal', id='A'))
        app.set_project(project)
        app.doctor_tree.selection_set('A')
        app.edit_weight()
        dialog = next(w for w in app.winfo_children() if isinstance(w, Form))
        dialog.values['coefficient'].delete(0, 'end')
        dialog.values['coefficient'].insert(0, value)
        with pytest.raises(ValueError, match='coefficient'):
            dialog.accept.__wrapped__(dialog)
        assert project.doctors()[0].coefficient == 1
    finally:
        app.close()


def test_equity_export_contains_the_weight_used_by_validated_plan(tmp_path, rules):
    doctors = [Doctor('Pondéré', id='A', coefficient=2), Doctor('Normal', id='B')]
    slots = weighted_slots()
    plan = solve(Plan(slots[0].day, slots[-1].day, slots), doctors, rules, [])
    path = tmp_path / 'Équité.xlsx'
    # Une version validée garde sa pondération même si la fiche est ensuite changée.
    with Project(tmp_path / 'Validé.gardavu', create=True) as project:
        project.save_rules(rules)
        for doctor in doctors:
            project.save_doctor(doctor)
        plan = project.validate_plan(plan)
        doctors[0].coefficient = 1
        project.save_doctor(doctors[0])
        export_excel(path, plan, project.doctors())
    workbook = load_workbook(path, data_only=False)
    try:
        eq = workbook['Équité']
        assert [eq.cell(1, col).value for col in range(1, 5)] == ['Médecin', 'Pondération', 'Cible', 'Attribué']
        assert [eq.cell(2, col).value for col in range(2, 5)] == [2, 6, 6]
        assert [eq.cell(3, col).value for col in range(2, 5)] == [1, 3, 3]
    finally:
        workbook.close()
