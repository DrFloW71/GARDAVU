from dataclasses import asdict
from datetime import date
import json
from pathlib import Path
import shutil

import pytest
from openpyxl import Workbook, load_workbook

from gardavu.domain import Doctor, Plan, Record, Slot, Span
from gardavu.exports import create_templates, export_csv, export_excel
from gardavu.imports import (Preview, commit_import, conflicts, prepare_doctors,
                             prepare_history, read_sheets)
from gardavu.storage import Project


def setup_project(path, rules):
    p = Project(path, create=True)
    p.save_doctor(Doctor('Été', 'Camille', id='A'))
    p.save_doctor(Doctor('Été', 'Camille', id='B'))
    p.save_doctor(Doctor('Autre', id='C'))
    p.save_rules(rules)
    return p


def ready_plan():
    s = Slot('2030-12-25:ferie-jour', '2030-12-25', 'FERIE_JOUR', '08:00', '19:00', holidays=('NOEL',))
    return Plan(s.day, s.day, [s], {s.id: ['A']})


def test_lock_backup_restore_reopen_and_accented_paths(tmp_path, rules):
    folder = tmp_path / 'Dossier partagé été'; folder.mkdir()
    path = folder / 'Gardes éç.gardavu'
    with setup_project(path, rules) as p:
        with pytest.raises(ValueError, match='déjà ouvert'):
            Project(path)
        p.save_draft(ready_plan())
        backup = p.backup()
        assert backup.exists()
        assert p.conn.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
    restored = folder / 'Restauration.gardavu'
    Project.restore(backup, restored)
    with Project(restored) as p:
        assert len(p.doctors()) == 3 and len(p.plans()) == 1
    other_pc = tmp_path / 'Copie autre ordinateur.gardavu'
    shutil.copy2(restored, other_pc)
    with Project(other_pc) as p:
        assert len(p.doctors()) == 3


def test_validation_versions_no_double_count_and_archive(tmp_path, rules):
    with setup_project(tmp_path / 'p.gardavu', rules) as p:
        original = ready_plan()
        validated = p.validate_plan(original)
        assert validated.state == 'VALIDE' and len(p.history()) == 1
        with pytest.raises(ValueError, match='immuable'):
            p.save_draft(original)
        second = ready_plan(); second.assignments[second.slots[0].id] = ['B']
        with pytest.raises(ValueError, match='Confirmez'):
            p.validate_plan(second)
        replaced = p.validate_plan(second, replace=True)
        assert len(p.history()) == 1 and p.history()[0].doctor_id == 'B'
        assert p.conn.execute('SELECT COUNT(*) FROM history').fetchone()[0] == 2
        assert any(v.state == 'REMPLACE' for v in p.plans())
        p.archive(replaced)
        assert len(p.history()) == 1 and replaced.state == 'ARCHIVE'


def test_holiday_repeat_requires_confirmation(tmp_path, rules):
    with setup_project(tmp_path / 'p.gardavu', rules) as p:
        old = Record(Slot('old', '2028-12-25', 'FERIE_NUIT', None, None, holidays=('NOEL',)), 'A')
        commit_import(p, Preview(records=[old]))
        with pytest.raises(ValueError, match='Confirmation explicite'):
            p.validate_plan(ready_plan())
        validated = p.validate_plan(ready_plan(), confirm_repeats=True)
        assert validated.state == 'VALIDE'


def test_independent_rejection_manual_swap(tmp_path, rules):
    with setup_project(tmp_path / 'p.gardavu', rules) as p:
        doctors = p.doctors()
        d = next(d for d in doctors if d.id == 'B')
        d.unavailable = [Span('2030-12-25', '2030-12-25')]; p.save_doctor(d)
        plan = ready_plan(); plan.assignments[plan.slots[0].id] = ['B']
        with pytest.raises(ValueError, match='indisponible'):
            p.validate_plan(plan)
        assert not p.history()


def test_import_names_missing_hours_and_idempotence(tmp_path, rules):
    with setup_project(tmp_path / 'p.gardavu', rules) as p:
        rows = [['Date', 'Médecin', 'Catégorie'], ['25/12/2028', 'ete camille', 'Férié jour']]
        cols = {'date': 'Date', 'medecin': 'Médecin', 'categorie': 'Catégorie'}
        preview = prepare_history(rows, cols, p.doctors(), rules)
        assert preview.unresolved == {'ete camille': ['A', 'B']}
        with pytest.raises(ValueError, match='rapprochements'):
            commit_import(p, preview)
        preview = prepare_history(rows, cols, p.doctors(), rules, {'ete camille': 'A'})
        assert preview.records[0].slot.start is None
        assert preview.records[0].slot.holidays == ('NOEL',)
        preview.fingerprint = 'file1'
        assert commit_import(p, preview)['imported'] == 1
        assert commit_import(p, preview)['already_imported']
        preview.fingerprint = 'differentfile'
        assert commit_import(p, preview)['imported'] == 0
        assert len(p.history()) == 1


def test_import_conflict_requires_choice_and_traceability(tmp_path, rules):
    with setup_project(tmp_path / 'p.gardavu', rules) as p:
        s = ready_plan().slots[0]
        commit_import(p, Preview(records=[Record(s, 'A')]))
        preview = Preview(records=[Record(s, 'B')])
        assert conflicts(p, preview)
        with pytest.raises(ValueError, match='explicitement'):
            commit_import(p, preview)
        assert commit_import(p, preview, 'ignore')['ignored'] == 1
        commit_import(p, preview, 'replace')
        assert len(p.history()) == 1 and p.history()[0].doctor_id == 'B'
        assert p.conn.execute('SELECT COUNT(*) FROM history').fetchone()[0] == 2


def test_multi_staff_reimport_preserves_unchanged_doctor(tmp_path, rules):
    from dataclasses import replace
    with setup_project(tmp_path / 'p.gardavu', rules) as p:
        s = replace(ready_plan().slots[0], required=2)
        commit_import(p, Preview(records=[Record(s, 'A'), Record(s, 'B')]))
        commit_import(p, Preview(records=[Record(s, 'A'), Record(s, 'C')]), 'replace')
        assert {r.doctor_id for r in p.history()} == {'A', 'C'}


def test_matrix_dates_duplicates_and_statuses(rules, doctors):
    rows = [['Date', 'Jour', 'Nuit'], [date(2026, 12, 25), 'DEMO-001', 'DEMO-002']]
    preview = prepare_history(rows, {'date': 'Date'}, doctors, rules, matrix={'Jour': 'FERIE_JOUR', 'Nuit': 'FERIE_NUIT'})
    assert len(preview.records) == 2 and not preview.errors
    rows = [['date', 'id', 'cat', 'status'], [date(2026, 12, 25), 'DEMO-001', 'FERIE_JOUR', 'BROUILLON']]
    preview = prepare_history(rows, {'date': 'date', 'id_medecin': 'id', 'categorie': 'cat', 'statut': 'status'}, doctors, rules)
    assert not preview.records and preview.warnings


def test_import_doctors_no_implicit_homonym_merge(rules):
    rows = [['Nom', 'Prénom', 'ID'], ['Été', 'Camille', None]]
    existing = [Doctor('Été', 'Camille', id='A')]
    preview = prepare_doctors(rows, {'nom': 'Nom', 'prenom': 'Prénom', 'id': 'ID'}, existing)
    assert preview.errors and not preview.doctors
    rows[1][2] = 'B'
    preview = prepare_doctors(rows, {'nom': 'Nom', 'prenom': 'Prénom', 'id': 'ID'}, existing)
    assert len(preview.doctors) == 1 and preview.doctors[0].id == 'B'


def test_formulas_rejected_and_exports_sanitized(tmp_path, doctors):
    wb = Workbook(); wb.active.append(['Date', 'Médecin']); wb.active.append(['=TODAY()', 'DEMO-001'])
    path = tmp_path / 'formula.xlsx'; wb.save(path)
    with pytest.raises(ValueError, match='Formule refusée'):
        read_sheets(path)
    doctors[0].display = '=HYPERLINK("bad")'
    p = ready_plan(); p.assignments[p.slots[0].id] = ['DEMO-001']
    export_excel(tmp_path / 'out.xlsx', p, doctors)
    wb = load_workbook(tmp_path / 'out.xlsx', data_only=False)
    assert wb.sheetnames == ['Planning', 'Équité', 'Jours fériés', 'Paramètres et avertissements']
    assert 'NON VALIDÉ' in wb['Planning']['A1'].value
    assert wb['Planning']['B4'].value.startswith("'")
    assert all(c.data_type != 'f' for ws in wb for row in ws for c in row)
    assert isinstance(wb['Planning']['A4'].value, __import__('datetime').datetime)
    export_csv(tmp_path / 'out.csv', p, doctors)
    assert "'=HYPERLINK" in (tmp_path / 'out.csv').read_text(encoding='utf-8-sig')


def test_templates_and_export_roundtrip(tmp_path, rules, doctors):
    assert len(create_templates(tmp_path)) == 3
    plan = ready_plan(); plan.state = 'VALIDE'; plan.assignments[plan.slots[0].id] = ['DEMO-001']
    path = tmp_path / 'planning.xlsx'; export_excel(path, plan, doctors)
    sheet = read_sheets(path)['Planning']
    mapping = {'date': 'Date', 'medecin': 'Médecin', 'id_medecin': 'ID médecin', 'categorie': 'Catégorie', 'debut': 'Début', 'fin': 'Fin', 'statut': 'Statut', 'fete': 'Fête', 'id_creneau': 'ID créneau', 'effectif': 'Effectif'}
    result = prepare_history(sheet, mapping, doctors, rules, header_row=2)
    assert not result.errors and len(result.records) == 1
    assert result.records[0].slot == plan.slots[0]


def test_coverage_not_inferred_from_endpoints(tmp_path, rules):
    with setup_project(tmp_path / 'p.gardavu', rules) as p:
        assert not p.coverage_complete('2030-01-01', 12)
        p.set_meta('coverage', [{'start': '2029-01-01', 'end': '2029-06-30'}, {'start': '2029-07-01', 'end': '2029-12-31'}])
        assert p.coverage_complete('2030-01-01', 12)


def test_transaction_rolls_back_and_future_schema_rejected(tmp_path, rules):
    path = tmp_path / 'p.gardavu'
    with setup_project(path, rules) as p:
        with pytest.raises(RuntimeError):
            with p.transaction():
                p.conn.execute("DELETE FROM doctors")
                raise RuntimeError('interruption')
        assert len(p.doctors()) == 3
        p.conn.execute('PRAGMA user_version=999')
    with pytest.raises(ValueError, match='plus récente'):
        Project(path)
